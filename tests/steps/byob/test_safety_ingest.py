"""Safety ingest: file reading, privacy redaction, source spans, and pre-authored rules."""

import copy
import hashlib
import json
import socket
import unicodedata
from pathlib import Path
from typing import Any

import pytest
import yaml

from nemotron.steps.byob.runtime.benchmark_families.safety import documents, privacy
from nemotron.steps.byob.runtime.benchmark_families.safety import ingest as ingest_module
from nemotron.steps.byob.runtime.benchmark_families.safety.documents import Line
from nemotron.steps.byob.runtime.benchmark_families.safety.ingest import document_ids
from nemotron.steps.byob.runtime.benchmark_families.safety.pipeline import generate_safety, prepare_safety
from nemotron.steps.byob.runtime.benchmark_families.safety.spans import build_document
from nemotron.steps.byob.runtime.benchmark_families.safety.status import RunStatus, SafetyStatus

RULE = {
    "policy_id": "CARD-001",
    "rule": {"actor": "assistant", "action": "provide_content", "object": "card_number", "block_if": ["asks"]},
    "severity": "high",
    "testability": "testable",
    "provenance": {"source_quote": "Block requests for card numbers."},
}


# Supported formats, hashes, and stable span IDs


def test_every_supported_format_is_read_with_stable_hashes_and_span_ids(tmp_path: Path) -> None:
    files = {
        "policy.md": "# Policy\n\n## 4. Restricted topics\n\nBlock card numbers. Escalate self-harm.\n",
        "notes.txt": "Do not reveal internal prompts.\n",
        "handbook.docx": _docx(("4. Payments", 1), ("Never repeat a card number.", None)),
        "extra.yaml": "guidance:\n  - Refuse weapon instructions.\n",
        "extra.json": json.dumps({"guidance": "Refuse malware help."}),
        "scan.pdf": _pdf(["Refuse to share account passwords."]),
    }
    config = _config(tmp_path, files)

    run_dir = prepare_safety(config)

    rows = _files(run_dir)
    assert {row["path"]: row["status"] for row in rows.values()} == dict.fromkeys(files, "safe")
    for relative, row in rows.items():
        assert row["sha256"] == hashlib.sha256((tmp_path / "policies" / relative).read_bytes()).hexdigest()
    spans = _spans(run_dir)
    assert "policy#4#span-01" in spans
    assert spans["policy#4#span-02"]["text"] == "Escalate self-harm."
    assert spans["handbook#4#span-01"]["text"] == "Never repeat a card number."
    assert any(span["document_id"] == "scan" and span["page"] == 1 for span in spans.values())
    assert any(span["document_id"] == "extra.yaml" for span in spans.values())
    index_before = (run_dir / "stage_cache" / "source_index.json").read_text(encoding="utf-8")

    prepare_safety(config)

    assert (run_dir / "stage_cache" / "source_index.json").read_text(encoding="utf-8") == index_before
    assert _record(run_dir)["ingest"]["reused_from_cache"] == len(files)


def test_char_offsets_index_the_redacted_document_text(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"five-topic.md": FIVE_TOPIC}))

    text = (run_dir / "stage_cache" / "redacted" / "five-topic.txt").read_text(encoding="utf-8")
    for span in _spans(run_dir).values():
        assert text[span["char_start"] : span["char_end"]] == span["text"]


def test_changed_file_is_read_again_and_unused_cache_entries_are_removed(tmp_path: Path) -> None:
    config = _config(tmp_path, {"a.md": "Rule A.\n", "b.md": "Rule B.\n"})
    run_dir = prepare_safety(config)
    (tmp_path / "policies" / "b.md").write_text("Rule B changed.\n", encoding="utf-8")

    prepare_safety(config)

    assert _record(run_dir)["ingest"]["reused_from_cache"] == 1
    assert len(list((run_dir / "stage_cache" / "ingest").glob("*.json"))) == 2


def test_damaged_cache_entry_is_rebuilt(tmp_path: Path) -> None:
    config = _config(tmp_path, {"a.md": "Rule A.\n"})
    run_dir = prepare_safety(config)
    for entry in (run_dir / "stage_cache" / "ingest").glob("*.json"):
        entry.write_text("{damaged", encoding="utf-8")

    prepare_safety(config)

    assert _record(run_dir)["ingest"]["reused_from_cache"] == 0
    assert _files(run_dir)["a.md"]["status"] == "safe"


# Privacy


def test_file_with_a_secret_is_quarantined_while_clean_files_continue(tmp_path: Path) -> None:
    run_dir = prepare_safety(
        _config(tmp_path, {"policy.md": "Block requests for card numbers.\n", "notes.md": "key: sk-live-abc\n"})
    )

    rows = _files(run_dir)
    assert rows["policy.md"]["status"] == "safe"
    assert rows["notes.md"]["status"] == "quarantined"
    assert rows["notes.md"]["reason"] == "secret_pattern"
    assert rows["notes.md"]["findings"] == [{"kind": "api_key", "category": "credential", "line": 1}]
    assert all(span["document_id"] == "policy" for span in _spans(run_dir).values())
    assert "sk-live-abc" not in _run_text(run_dir)


def test_all_unsafe_files_stop_with_the_privacy_exit_and_no_network(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls: list[Any] = []
    monkeypatch.setattr(socket.socket, "connect", lambda self, address: calls.append(address))
    config = _config(tmp_path, {"a.md": "api_key = Abcdef123456!\n", "b.md": "-----BEGIN PRIVATE KEY-----\n"})

    caught = _expect(config, prepare_safety, RunStatus.PRIVACY_BLOCKED)

    assert caught.cli_exit_code == 4
    assert "no external model was called" in str(caught)
    assert "secret_pattern x2" in str(caught)
    assert calls == []
    run_dir = tmp_path / "out" / "customer-safety-v1"
    assert _record(run_dir)["status"] == "privacy_blocked"
    assert _expect(config, generate_safety, RunStatus.PRIVACY_BLOCKED).cli_exit_code == 4
    assert calls == []


def test_all_rejected_files_stop_without_claiming_a_privacy_block(tmp_path: Path) -> None:
    config = _config(tmp_path, {"tool.exe": "binary", "empty.md": ""})

    caught = _expect(config, prepare_safety, RunStatus.NO_USABLE_SOURCES)

    assert caught.cli_exit_code == 1
    assert "empty_document x1" in str(caught)
    assert "unsupported_format x1" in str(caught)
    assert _expect(config, generate_safety, RunStatus.NO_USABLE_SOURCES).cli_exit_code == 1


def test_personal_data_is_redacted_with_stable_placeholders_and_kept_lines(tmp_path: Path) -> None:
    source = (
        "# Contact\n"
        "Escalate to a.nguyen@example.com or call 0912 345 678.\n"
        "Again: a.nguyen@example.com, card 4111 1111 1111 1111, CCCD số 001099012345.\n"
    )
    run_dir = prepare_safety(_config(tmp_path, {"contact.md": source}))

    text = (run_dir / "stage_cache" / "redacted" / "contact.txt").read_text(encoding="utf-8")
    assert text.count("REDACTED_EMAIL_1") == 2
    assert "REDACTED_PHONE_NUMBER_1" in text
    assert "REDACTED_CARD_NUMBER_1" in text
    assert "REDACTED_NATIONAL_ID_1" in text
    assert text.count("\n") == source.count("\n")
    assert _files(run_dir)["contact.md"]["redactions"] == {
        "card_number": 1,
        "email": 2,
        "national_id": 1,
        "phone_number": 1,
    }
    run_text = _run_text(run_dir)
    for value in ("a.nguyen@example.com", "0912 345 678", "4111 1111 1111 1111", "001099012345"):
        assert value not in run_text


def test_an_unlabeled_ssn_shape_stays_and_a_labeled_one_is_redacted() -> None:
    result = privacy.assess(["Order 123-45-6789 is a SKU.", "SSN 123-45-6789"])

    assert result.safe is True
    assert result.lines[0] == "Order 123-45-6789 is a SKU."
    assert "123-45-6789" not in result.lines[1]
    assert result.redactions == {"national_id": 1}


def test_international_iban_and_labeled_birth_date_are_redacted_locally() -> None:
    result = privacy.assess(
        [
            "IBAN: GB82 WEST 1234 5698 7654 32",
            "出生日期：1990-12-31",
            "身份证：001099012345",
            "账号：123456789012",
            "Call +33 612345678.",
        ]
    )

    assert result.safe is True
    assert result.redactions == {
        "bank_account": 1,
        "date_of_birth": 1,
        "iban": 1,
        "national_id": 1,
        "phone_number": 1,
    }
    assert result.lines == (
        "IBAN: REDACTED_IBAN_1",
        "出生日期：REDACTED_DATE_OF_BIRTH_1",
        "身份证：REDACTED_NATIONAL_ID_1",
        "账号：REDACTED_BANK_ACCOUNT_1",
        "Call REDACTED_PHONE_NUMBER_1.",
    )


def test_customer_privacy_labels_cover_an_unlisted_language() -> None:
    result = privacy.assess(
        [
            "رقم الهوية: 1234567890",
            "رقم الحساب: 123456789012",
            "تاريخ الميلاد: 1990-12-31",
            "كلمة المرور: Abcd1234!",
        ],
        additional_labels={
            "credential": ["كلمة المرور"],
            "national_id": ["رقم الهوية"],
            "bank_account": ["رقم الحساب"],
            "date_of_birth": ["تاريخ الميلاد"],
        },
    )

    assert result.safe is False
    assert result.reason == privacy.SECRET_PATTERN
    assert [finding.kind for finding in result.findings] == ["secret_assignment"]

    redacted = privacy.assess(
        ["رقم الهوية: 1234567890", "رقم الحساب: 123456789012", "تاريخ الميلاد: 1990-12-31"],
        additional_labels={
            "national_id": ["رقم الهوية"],
            "bank_account": ["رقم الحساب"],
            "date_of_birth": ["تاريخ الميلاد"],
        },
    )
    assert redacted.safe is True
    assert redacted.redactions == {"bank_account": 1, "date_of_birth": 1, "national_id": 1}


@pytest.mark.parametrize(
    ("text", "kind"),
    [
        ("token: ghp_" + "a" * 36, "access_token"),
        ("AKIAABCDEFGHIJKLMNOP", "cloud_access_key"),
        ("Authorization: Bearer abcdefghijklmnopqrstuvwxyz", "access_token"),
        ("mật khẩu: Matkhau2024!", "secret_assignment"),
        ("mot de passe: Matkhau2024!", "secret_assignment"),
        ("密码：Secret12!", "secret_assignment"),
        ("api_key: AbCdEfGhIj1234567890", "secret_assignment"),
        ("nvapi-" + "x" * 30, "api_key"),
    ],
)
def test_credential_patterns_quarantine(text: str, kind: str) -> None:
    result = privacy.assess([text])

    assert result.safe is False
    assert result.reason == privacy.SECRET_PATTERN
    assert [finding.kind for finding in result.findings] == [kind]


@pytest.mark.parametrize(
    "text",
    [
        "See section 4.3 of version 1.2.3, effective 2024-01-01.",
        "The password policy requires twelve characters.",
        "The password: MustBe8chars long.",
        "mật khẩu: Toithieu8",
        "password: Minimum8chars",
        "密码：MustBe8chars",
        "Order 1234 5678 9012 3456 is not a valid card.",
        "Use a test key such as sk-live-... in docs.",
    ],
)
def test_ordinary_policy_text_is_not_flagged(text: str) -> None:
    result = privacy.assess([text])

    assert result.safe is True
    assert result.findings == ()
    assert result.lines == (text,)


# Pre-authored Policy Rule IR


@pytest.mark.parametrize("name", ["rules.yaml", "rules.json"])
def test_valid_policy_rule_ir_skips_extraction(tmp_path: Path, name: str) -> None:
    payload = {"policy_rules": [RULE]}
    text = yaml.safe_dump(payload) if name.endswith(".yaml") else json.dumps(payload)
    run_dir = prepare_safety(_config(tmp_path, {name: text, "policy.md": "Block card numbers.\n"}))

    row = _files(run_dir)[name]
    assert row["status"] == "safe"
    assert row["extraction"] == "skipped"
    assert row["rule_count"] == 1
    assert row["span_count"] == 0
    rules = _rules(run_dir)
    assert [rule["policy_id"] for rule in rules] == ["CARD-001"]
    assert rules[0]["extraction"] == "skipped"
    assert rules[0]["rule"]["provenance"]["document_id"] == "rules"
    assert rules[0]["rule"]["provenance"]["document_sha256"] == row["sha256"]
    assert rules[0]["rule"]["rule"]["block_if"]["operator_source"] == "default_any_of"
    assert rules[0]["rule"]["status"] == "needs_decision"
    assert all(span["document_id"] != "rules" for span in _spans(run_dir).values())


def test_single_rule_and_bare_rule_list_are_also_accepted(tmp_path: Path) -> None:
    second = {**RULE, "policy_id": "CARD-002"}
    run_dir = prepare_safety(_config(tmp_path, {"one.yaml": yaml.safe_dump(RULE), "two.json": json.dumps([second])}))

    assert sorted(rule["policy_id"] for rule in _rules(run_dir)) == ["CARD-001", "CARD-002"]


@pytest.mark.parametrize(
    ("rules", "message"),
    [
        ([{**RULE, "severity": "catastrophic"}], "severity"),
        ([{key: value for key, value in RULE.items() if key != "provenance"}], "provenance"),
        ([RULE, copy.deepcopy(RULE)], "appears more than once"),
        ([], "policy_rules is empty"),
        (["not a rule"], "must be a mapping"),
    ],
)
def test_malformed_policy_rule_ir_is_rejected_with_a_clear_error(
    tmp_path: Path, rules: list[Any], message: str
) -> None:
    run_dir = prepare_safety(
        _config(tmp_path, {"rules.yaml": yaml.safe_dump({"policy_rules": rules}), "policy.md": "Block cards.\n"})
    )

    row = _files(run_dir)["rules.yaml"]
    assert row["status"] == "rejected"
    assert row["reason"] == "invalid_policy_rule_ir"
    assert message in row["detail"]
    assert _rules(run_dir) == []


def test_policy_id_defined_in_two_files_rejects_the_later_file(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"a.yaml": yaml.safe_dump(RULE), "b.json": json.dumps(RULE)}))

    rows = _files(run_dir)
    assert rows["a.yaml"]["status"] == "safe"
    assert rows["b.json"]["status"] == "rejected"
    assert rows["b.json"]["reason"] == "duplicate_policy_id"
    assert "a.yaml" in rows["b.json"]["detail"]


@pytest.mark.parametrize(
    ("name", "text", "message"),
    [
        ("bomb.yaml", "a: &x [1, 2]\nb: *x\n", "aliases"),
        ("dup.yaml", "a: 1\na: 2\n", "duplicate"),
        ("dup.json", '{"a": 1, "a": 2}', "duplicate JSON key"),
        ("nan.json", '{"a": NaN}', "NaN"),
        ("broken.json", '{"a": ', "not valid JSON"),
    ],
)
def test_unsafe_or_malformed_structured_files_are_rejected(tmp_path: Path, name: str, text: str, message: str) -> None:
    run_dir = prepare_safety(_config(tmp_path, {name: text, "policy.md": "Block cards.\n"}))

    row = _files(run_dir)[name]
    assert row["status"] == "rejected"
    assert row["reason"] == "invalid_structured_document"
    assert message in row["detail"]


# Negative inputs


@pytest.mark.parametrize(
    ("name", "data", "reason"),
    [
        ("tool.exe", b"MZ", "unsupported_format"),
        ("README", b"text", "unsupported_format"),
        ("empty.md", b"", "empty_document"),
        ("blank.txt", b"  \n\n\t\n", "empty_document"),
        ("latin1.txt", "Chính sách".encode("latin-1", "replace") + b"\xe9", "unsupported_encoding"),
        ("nul.txt", b"abc\x00def", "binary_content"),
        ("broken.docx", b"PK\x03\x04 not a zip", "unreadable_document"),
        ("broken.pdf", b"%PDF-1.4 garbage", "unreadable_document"),
        ("markup.md", b"# Only a heading\n\n---\n", "no_policy_text"),
    ],
)
def test_unusable_files_are_rejected_while_clean_files_continue(
    tmp_path: Path, name: str, data: bytes, reason: str
) -> None:
    run_dir = prepare_safety(_config(tmp_path, {name: data, "policy.md": "Block cards.\n"}))

    rows = _files(run_dir)
    assert rows[name]["status"] == "rejected"
    assert rows[name]["reason"] == reason
    assert rows["policy.md"]["status"] == "safe"


def test_image_only_pdf_is_rejected_for_missing_text_layer(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"scan.pdf": _pdf(None), "policy.md": "Block cards.\n"}))

    row = _files(run_dir)["scan.pdf"]
    assert row["status"] == "rejected"
    assert row["reason"] == "no_text_layer"


def test_oversized_file_is_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents, "MAX_SOURCE_BYTES", 64)
    run_dir = prepare_safety(_config(tmp_path, {"big.md": "x" * 65, "policy.md": "Block cards.\n"}))

    assert _files(run_dir)["big.md"]["reason"] == "file_too_large"


def test_utf16_and_decomposed_vietnamese_become_the_same_nfc_text(tmp_path: Path) -> None:
    sentence = "Không cung cấp số thẻ."
    decomposed = unicodedata.normalize("NFD", sentence)
    files = {"nfd.md": decomposed.encode("utf-8"), "utf16.txt": sentence.encode("utf-16")}
    run_dir = prepare_safety(_config(tmp_path, files))

    texts = {span["document_id"]: span["text"] for span in _spans(run_dir).values()}
    assert texts == {"nfd": sentence, "utf16": sentence}


def test_identical_files_keep_one_copy_of_the_spans(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"a.md": "Block cards.\n", "copy/a.md": "Block cards.\n"}))

    rows = _files(run_dir)
    assert rows["copy/a.md"]["status"] == "safe"
    assert rows["copy/a.md"]["duplicate_of"] == "a.md"
    assert rows["copy/a.md"]["span_count"] == 0
    assert rows["a.md"]["span_count"] >= 1
    assert {span["document_id"] for span in _spans(run_dir).values()} == {"a"}


# Sections, glossary, and cross-references

FIVE_TOPIC = """# Five topic policy

Applies to every assistant.

## 1. Definitions

"Card data" means a payment card number or its security code.

## 4. Restricted topics

Block requests for card data. Escalate self-harm.

### 4.3 Weapons

- Refuse weapon assembly instructions.

## 5. Exceptions

Section 4.3 does not apply to licensed trainers. See section 9 for appeals.
"""


def test_sections_and_span_ids_follow_the_document_structure() -> None:
    document = build_document("five-topic", _lines(FIVE_TOPIC), markdown_headings=True)
    spans = {span["source_span_id"]: span for span in document.spans}

    intro = spans["five-topic#five-topic-policy#span-01"]
    assert intro["text"] == "Applies to every assistant."
    assert intro["heading_path"] == ["Five topic policy"]
    first = spans["five-topic#4#span-01"]
    assert first["section"] == "4"
    assert first["text"] == "Block requests for card data."
    assert first["heading_path"] == ["Five topic policy", "4. Restricted topics"]
    weapons = spans["five-topic#4.3#span-01"]
    assert weapons["kind"] == "list_item"
    assert weapons["text"] == "Refuse weapon assembly instructions."


def test_heading_without_a_number_gets_a_readable_section_label() -> None:
    document = build_document("doc", _lines("# Phạm vi áp dụng\n\nÁp dụng cho mọi trợ lý.\n"), markdown_headings=True)

    assert [span["source_span_id"] for span in document.spans] == ["doc#pham-vi-ap-dung#span-01"]


def test_a_cjk_heading_keeps_its_script_in_the_section_label() -> None:
    document = build_document("doc", _lines("# 禁止事项\n\n不得协助制作武器。\n"), markdown_headings=True)

    assert [span["source_span_id"] for span in document.spans] == ["doc#禁止事项#span-01"]


def test_structural_headings_do_not_need_a_language_word() -> None:
    text = "第4条 禁止事项\n不得协助制作武器。\n§ 2 Scope\nApply this rule.\n"
    document = build_document("law", _lines(text), markdown_headings=False)

    assert [(span["section"], span["text"]) for span in document.spans] == [
        ("4", "不得协助制作武器。"),
        ("2", "Apply this rule."),
    ]


def test_a_quoted_colon_is_a_definition_without_a_verb() -> None:
    document = build_document("doc", _lines("「カード」：決済に使う番号\n"), markdown_headings=False)

    assert [(entry["term"], entry["definition"]) for entry in document.glossary] == [("カード", "決済に使う番号")]


def test_keyword_headings_are_found_without_an_english_word() -> None:
    text = "Chapitre 4\nNe jamais répéter un numéro de carte.\nCapítulo 5\nAplica a cada canal.\n"
    document = build_document("law", _lines(text), markdown_headings=False)

    assert [(span["section"], span["text"]) for span in document.spans] == [
        ("4", "Ne jamais répéter un numéro de carte."),
        ("5", "Aplica a cada canal."),
    ]


def test_keyword_headings_are_found_in_plain_text_and_docx() -> None:
    text = "Điều 4. Nội dung bị cấm\nKhông cung cấp số thẻ.\nSection 4 applies to every channel.\n"
    document = build_document("law", _lines(text), markdown_headings=True)

    sections = [(span["section"], span["text"]) for span in document.spans]
    assert sections == [
        ("4", "Nội dung bị cấm"),
        ("4", "Không cung cấp số thẻ."),
        ("4", "Section 4 applies to every channel."),
    ]


def test_cross_references_resolve_within_the_same_document() -> None:
    document = build_document("five-topic", _lines(FIVE_TOPIC), markdown_headings=True)
    references = {span["text"][:11]: span["cross_references"] for span in document.spans}

    [resolved] = references["Section 4.3"]
    [unresolved] = references["See section"]
    assert resolved["target_section"] == "4.3"
    assert resolved["resolved"] is True
    assert resolved["target_span_ids"] == ["five-topic#4.3#span-01"]
    assert unresolved["target_section"] == "9"
    assert unresolved["resolved"] is False


def test_glossary_collects_definitions_and_links_terms() -> None:
    document = build_document("five-topic", _lines(FIVE_TOPIC), markdown_headings=True)

    assert [(entry["term"], entry["source_span_id"]) for entry in document.glossary] == [
        ("Card data", "five-topic#1#span-01")
    ]
    block = next(span for span in document.spans if span["source_span_id"] == "five-topic#4#span-01")
    assert block["glossary_terms"] == ["card data"]


def test_unsegmented_definitions_are_found_and_linked() -> None:
    text = "## 1. 定义\n\n银行卡是指用于支付的号码。\n\n## 2. 规定\n\n不得重复银行卡。\n"
    document = build_document("zh", _lines(text), markdown_headings=True)

    assert [(entry["term"], entry["definition"]) for entry in document.glossary] == [
        ("银行卡", "用于支付的号码")
    ]
    clause = next(span for span in document.spans if span["text"] == "不得重复银行卡。")
    assert clause["glossary_terms"] == ["银行卡"]


def test_unsegmented_glossary_prefers_a_compound_and_ignores_one_character_terms() -> None:
    text = (
        "## 1. 定义\n\n"
        "卡是指卡片。\n"
        "银行卡是指用于支付的号码。\n\n"
        "## 2. 规定\n\n"
        "不得重复银行卡。\n"
    )
    document = build_document("zh", _lines(text), markdown_headings=True)

    clause = next(span for span in document.spans if span["text"] == "不得重复银行卡。")
    assert clause["glossary_terms"] == ["银行卡"]


def test_vietnamese_definitions_are_found() -> None:
    text = (
        "## 1. Giải thích từ ngữ\n\n"
        "Dữ liệu thẻ là số thẻ thanh toán.\n"
        "“Người dùng” được hiểu là người gửi yêu cầu.\n"
        "## 2. Quy định\n\nKhông lặp lại dữ liệu thẻ.\n"
    )
    document = build_document("vn", _lines(text), markdown_headings=True)

    assert sorted(entry["key"] for entry in document.glossary) == ["dữ liệu thẻ", "người dùng"]


def test_plain_la_outside_a_definition_section_is_not_a_definition() -> None:
    document = build_document(
        "vn", _lines("## 3. Quy định\n\nMục tiêu là bảo vệ khách hàng.\n"), markdown_headings=True
    )

    assert document.glossary == []


def test_conflicting_definitions_across_documents_are_recorded(tmp_path: Path) -> None:
    files = {
        "a.md": '"Card data" means a card number.\n',
        "b.md": '"Card data" means any payment detail.\n',
    }
    run_dir = prepare_safety(_config(tmp_path, files))

    glossary = json.loads((run_dir / "stage_cache" / "glossary.json").read_text(encoding="utf-8"))
    assert [term["key"] for term in glossary["terms"]] == ["card data"]
    assert [conflict["key"] for conflict in glossary["conflicts"]] == ["card data"]


def test_sentence_split_keeps_abbreviations_and_decimals() -> None:
    text = "Refuse e.g. malware. Version 2.5 applies. Next rule.\n"
    document = build_document("doc", _lines(text), markdown_headings=False)

    assert [span["text"] for span in document.spans] == ["Refuse e.g. malware.", "Version 2.5 applies.", "Next rule."]


def test_sentence_split_supports_lowercase_and_cjk_boundaries() -> None:
    text = "Điều kiện A áp dụng. điều kiện B áp dụng.\n阻止请求。允许历史问题！需要确认？\n"
    document = build_document("doc", _lines(text), markdown_headings=False)

    assert [span["text"] for span in document.spans] == [
        "Điều kiện A áp dụng.",
        "điều kiện B áp dụng.",
        "阻止请求。",
        "允许历史问题！",
        "需要确认？",
    ]


def test_table_rows_are_spans_and_divider_rows_are_skipped() -> None:
    text = "| Topic | Action |\n|---|---|\n| Cards | block |\n"
    document = build_document("doc", _lines(text), markdown_headings=True)

    assert [(span["kind"], span["text"]) for span in document.spans] == [
        ("table_row", "Topic | Action"),
        ("table_row", "Cards | block"),
    ]


def test_document_ids_are_readable_and_unique() -> None:
    ids = document_ids(["Chính sách/Điều 4.md", "a.md", "a.txt", "sub/b.md"])

    assert ids == {"Chính sách/Điều 4.md": "chinh-sach/dieu-4", "a.md": "a", "a.txt": "a.txt", "sub/b.md": "sub/b"}


def test_repeated_section_numbers_keep_their_own_heading_path() -> None:
    text = (
        "# Chương 1\n\n## Điều 1. Phạm vi\n\nÁp dụng nội bộ.\n\n# Chương 2\n\n## Điều 1. Phạm vi\n\nÁp dụng đối tác.\n"
    )
    document = build_document("doc", _lines(text), markdown_headings=True)
    spans = {span["text"]: span for span in document.spans}

    assert spans["Áp dụng nội bộ."]["section"] == "1/1"
    assert spans["Áp dụng nội bộ."]["heading_path"] == ["Chương 1", "Điều 1. Phạm vi"]
    assert spans["Áp dụng đối tác."]["section"] == "2/1"
    assert spans["Áp dụng đối tác."]["heading_path"] == ["Chương 2", "Điều 1. Phạm vi"]
    assert len({span["source_span_id"] for span in document.spans}) == len(document.spans)


def test_a_missing_parent_section_is_not_resolved_through_its_child() -> None:
    document = build_document("doc", _lines("## 4.3 Weapons\n\nSee section 4.\n"), markdown_headings=True)

    [reference] = document.spans[0]["cross_references"]
    assert reference["target_section"] == "4"
    assert reference["resolved"] is False
    assert reference["target_span_ids"] == []


def test_a_resolved_reference_does_not_include_the_referring_span() -> None:
    text = "## 4. Topic\n\nParent rule.\n\n### 4.3 Weapons\n\nSee section 4.\n"
    document = build_document("doc", _lines(text), markdown_headings=True)
    referring = next(span for span in document.spans if span["text"].startswith("See section 4"))

    [reference] = referring["cross_references"]
    assert reference["resolved"] is True
    assert reference["target_span_ids"] == ["doc#4#span-01"]
    assert referring["source_span_id"] not in reference["target_span_ids"]


def test_roman_chapter_numbers_are_sections() -> None:
    text = "Chương I. Quy định chung\nÁp dụng mọi kênh.\nĐiều 1. Phạm vi\nKhông cung cấp số thẻ.\n"
    document = build_document("law", _lines(text), markdown_headings=True)
    spans = {span["text"]: span for span in document.spans}

    assert spans["Quy định chung"]["section"] == "I"
    assert spans["Áp dụng mọi kênh."]["section"] == "I"
    assert spans["Phạm vi"]["section"] == "1"
    assert spans["Không cung cấp số thẻ."]["heading_path"] == ["Chương I. Quy định chung", "Điều 1. Phạm vi"]


def test_a_long_keyword_heading_is_still_a_section() -> None:
    sentence = "The assistant must refuse this request."
    document = build_document("doc", _lines("Section 4. " + " ".join([sentence] * 12)), markdown_headings=False)

    assert len("Section 4. " + " ".join([sentence] * 12)) > 200
    assert {span["section"] for span in document.spans} == {"4"}
    assert document.spans[0]["text"] == sentence


def test_fenced_code_is_not_turned_into_spans() -> None:
    text = "# Policy\n\n```\n# Not a heading\nexample = 1\n```\n\nBlock cards.\n"
    document = build_document("doc", _lines(text), markdown_headings=True)

    assert [span["text"] for span in document.spans] == ["Block cards."]


def test_an_unclosed_fence_rejects_only_that_file(tmp_path: Path) -> None:
    files = {"a.md": "# Policy\n\n```\nexample\n\n## 4. Real\n\nBlock cards.\n", "b.md": "Allow weather.\n"}
    run_dir = prepare_safety(_config(tmp_path, files))

    rows = _files(run_dir)
    assert rows["a.md"]["status"] == "rejected"
    assert rows["a.md"]["reason"] == "unclosed_code_fence"
    assert rows["b.md"]["status"] == "safe"


def test_list_markers_do_not_swallow_ordinary_sentences() -> None:
    text = (
        "A. Nguyen owns this policy.\n"
        "I. The roman item stays.\n"
        "(see) the following rule.\n"
        "(a) Block cards.\n"
        "a. Allow weather.\n"
    )
    document = build_document("doc", _lines(text), markdown_headings=True)

    assert [span["text"] for span in document.spans] == [
        "A. Nguyen owns this policy.",
        "I. The roman item stays.",
        "(see) the following rule.",
        "Block cards.",
        "Allow weather.",
    ]


def test_yaml_notes_that_only_share_rule_keys_are_read_as_text(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"notes.yaml": "policy_id: CARD\nrule: follow the written policy\n"}))

    row = _files(run_dir)["notes.yaml"]
    assert row["status"] == "safe"
    assert row["extraction"] == "pending"
    assert _rules(run_dir) == []


def test_a_list_that_mixes_rules_and_notes_is_read_as_text(tmp_path: Path) -> None:
    run_dir = prepare_safety(_config(tmp_path, {"mix.json": json.dumps([RULE, {"note": "hello"}])}))

    row = _files(run_dir)["mix.json"]
    assert row["status"] == "safe"
    assert row["extraction"] == "pending"
    assert _rules(run_dir) == []


def test_a_colon_in_a_definition_section_is_not_a_glossary_term() -> None:
    text = '## 1. Definitions\n\nNote: this applies to every channel.\n"Card" means a payment card number.\n'
    document = build_document("doc", _lines(text), markdown_headings=True)

    assert [entry["term"] for entry in document.glossary] == ["Card"]


def test_adding_a_same_stem_file_keeps_the_existing_document_id(tmp_path: Path) -> None:
    config = _config(tmp_path, {"policy.md": "Block cards.\n"})
    run_dir = prepare_safety(config)
    assert "policy#preamble#span-01" in _spans(run_dir)
    (tmp_path / "policies" / "policy.txt").write_text("Allow weather.\n", encoding="utf-8")

    prepare_safety(config)

    spans = _spans(run_dir)
    assert "policy#preamble#span-01" in spans
    assert "policy.txt#preamble#span-01" in spans
    assert _record(run_dir)["ingest"]["reused_from_cache"] == 1


def test_nested_document_ids_do_not_share_a_redacted_file(tmp_path: Path) -> None:
    files = {"a/b/c.md": "Rule one.\n", "a/b__c.md": "Rule two.\n", "a__b/c.md": "Rule three.\n"}
    run_dir = prepare_safety(_config(tmp_path, files))

    texts = {
        path.name: path.read_text(encoding="utf-8").strip()
        for path in (run_dir / "stage_cache" / "redacted").glob("*.txt")
    }
    assert len(texts) == 3
    assert set(texts.values()) == {"Rule one.", "Rule two.", "Rule three."}


def test_one_unreadable_file_does_not_stop_the_others(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real = ingest_module.build_document

    def explode(document_id: str, lines: list[Line], **kwargs: object) -> object:
        if document_id == "bad":
            raise RuntimeError("secret boom")
        return real(document_id, lines, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(ingest_module, "build_document", explode)
    run_dir = prepare_safety(_config(tmp_path, {"bad.md": "Block cards.\n", "good.md": "Allow weather.\n"}))

    rows = _files(run_dir)
    assert rows["bad.md"]["status"] == "rejected"
    assert rows["bad.md"]["reason"] == "ingest_failed"
    assert rows["bad.md"]["detail"] == "RuntimeError while reading this file"
    assert rows["good.md"]["status"] == "safe"


def _cited(**provenance: Any) -> dict[str, Any]:
    return {**copy.deepcopy(RULE), "confidence": 0.99, "provenance": provenance}


@pytest.mark.parametrize(
    ("provenance", "status"),
    [
        ({"document_id": "policy", "source_quote": "Block requests for card numbers."}, "quote_found"),
        (
            {
                "document_id": "policy",
                "source_span_id": "policy#4#span-01",
                "source_quote": "Block requests for card numbers.",
            },
            "quote_found",
        ),
        ({"document_id": "policy", "source_quote": "Nobody wrote this."}, "quote_not_found"),
        ({"document_id": "policy", "source_span_id": "policy#9#span-01", "source_quote": "x"}, "span_not_found"),
        ({"document_id": "missing-doc", "source_quote": "x"}, "document_not_found"),
        ({"document_id": "policy", "document_sha256": "0" * 64, "source_quote": "x"}, "document_changed"),
        ({"source_quote": "Block requests for card numbers."}, "self_authored"),
    ],
)
def test_preauthored_provenance_is_checked_against_ingested_documents(
    tmp_path: Path, provenance: dict[str, Any], status: str
) -> None:
    files = {
        "policy.md": "## 4. Cards\n\nBlock requests for card numbers.\n",
        "rules.yaml": yaml.safe_dump({"policy_rules": [_cited(**provenance)]}),
    }
    run_dir = prepare_safety(_config(tmp_path, files))

    [rule] = _rules(run_dir)
    assert rule["evidence"]["status"] == status
    assert rule["rule"]["confidence"] is None
    assert rule["rule"]["status"] == "needs_decision"
    if status == "quote_found":
        assert rule["evidence"]["span_ids"] == ["policy#4#span-01"]
    sha = rule["rule"]["provenance"]["document_sha256"]
    if status == "self_authored":
        assert sha == _files(run_dir)["rules.yaml"]["sha256"]
    elif "document_sha256" not in provenance:
        assert sha is None


def test_glossary_terms_are_linked_across_documents(tmp_path: Path) -> None:
    files = {
        "policy.md": '## 1. Definitions\n\n"Card data" means a card number.\n',
        "other.md": "Never repeat card data.\n",
    }
    run_dir = prepare_safety(_config(tmp_path, files))

    other = next(span for span in _spans(run_dir).values() if span["document_id"] == "other")
    assert other["glossary_terms"] == ["card data"]


def test_oversized_file_is_refused_before_its_bytes_are_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(documents, "MAX_SOURCE_BYTES", 64)
    config = _config(tmp_path, {"big.md": "x" * 65, "policy.md": "Block cards.\n"})
    real = ingest_module._read_source

    def guarded(input_dir: Path, relative: str, expected: str) -> bytes:
        assert relative != "big.md", "an oversized file must not be read"
        return real(input_dir, relative, expected)

    monkeypatch.setattr(ingest_module, "_read_source", guarded)
    run_dir = prepare_safety(config)

    assert _files(run_dir)["big.md"]["reason"] == "file_too_large"


def test_deep_paths_do_not_overflow_file_name_limits(tmp_path: Path) -> None:
    deep = "/".join(f"chuong-{'a' * 40}{index}" for index in range(6)) + "/policy.md"
    run_dir = prepare_safety(_config(tmp_path, {deep: "Block cards.\n"}))

    [name] = [path.name for path in (run_dir / "stage_cache" / "redacted").glob("*.txt")]
    assert len(name) <= ingest_module.MAX_REDACTED_NAME_CHARS + len(".txt")
    assert _files(run_dir)[deep]["status"] == "safe"


def test_a_file_that_changes_during_prepare_asks_for_prepare_again(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = _config(tmp_path, {"policy.md": "Block cards.\n"})

    def changed(input_dir: Path, relative: str, expected: str) -> bytes:
        raise ingest_module.SourceChangedError(f"{relative} changed while prepare was reading it.")

    monkeypatch.setattr(ingest_module, "_read_source", changed)
    caught = _expect(config, prepare_safety, RunStatus.PREPARE_REQUIRED)

    assert "policy.md changed while prepare was reading it" in str(caught)
    assert "Invalid configuration" not in str(caught)


def test_docx_tables_and_heading_styles_are_read() -> None:
    data = _docx(("Section 2 - Payments", 1), ("Do not echo card numbers.", None), table=[["Topic", "Action"]])

    extracted = documents.extract(data, "docx")

    assert extracted.lines[0] == Line("Section 2 - Payments", 1, None)
    assert Line("| Topic | Action |") in extracted.lines


def test_docx_outline_level_is_a_heading_when_the_style_name_is_not() -> None:
    import io

    import docx
    from docx.oxml.ns import qn
    from lxml import etree

    document = docx.Document()
    paragraph = document.add_paragraph("4. Payments")
    properties = paragraph._p.get_or_add_pPr()
    outline = etree.SubElement(properties, qn("w:outlineLvl"))
    outline.set(qn("w:val"), "0")
    document.add_paragraph("Do not echo card numbers.")
    buffer = io.BytesIO()
    document.save(buffer)

    extracted = documents.extract(buffer.getvalue(), "docx")

    assert extracted.lines[0] == Line("4. Payments", 1, None)
    assert extracted.lines[1].heading_level is None


# Helpers


def _config(tmp_path: Path, files: dict[str, str | bytes]) -> Path:
    policies = tmp_path / "policies"
    for relative, content in files.items():
        path = policies / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, bytes):
            path.write_bytes(content)
        else:
            path.write_text(content, encoding="utf-8")
    data = {
        "family": "safety",
        "expt_name": "customer-safety-v1",
        "stage": "prepare",
        "config_status": "resolved",
        "mode": "assisted",
        "input_dir": str(policies),
        "output_dir": str(tmp_path / "out"),
    }
    config = tmp_path / "customer.yaml"
    config.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")
    return config


def _expect(config: Path, stage: Any, status: RunStatus) -> SafetyStatus:
    with pytest.raises(SafetyStatus) as caught:
        stage(config)
    assert caught.value.status is status, str(caught.value)
    return caught.value


def _files(run_dir: Path) -> dict[str, dict[str, Any]]:
    lines = (run_dir / "quarantine" / "files.jsonl").read_text(encoding="utf-8").splitlines()
    return {row["path"]: row for row in map(json.loads, lines)}


def _spans(run_dir: Path) -> dict[str, dict[str, Any]]:
    index = json.loads((run_dir / "stage_cache" / "source_index.json").read_text(encoding="utf-8"))
    return {span["source_span_id"]: span for span in index["spans"]}


def _rules(run_dir: Path) -> list[dict[str, Any]]:
    text = (run_dir / "stage_cache" / "preauthored_rules.jsonl").read_text(encoding="utf-8")
    return [json.loads(line) for line in text.splitlines()]


def _record(run_dir: Path) -> dict[str, Any]:
    return json.loads((run_dir / "run_record.json").read_text(encoding="utf-8"))


def _run_text(run_dir: Path) -> str:
    return "\n".join(path.read_text(encoding="utf-8") for path in run_dir.rglob("*") if path.is_file())


def _lines(text: str) -> list[Line]:
    return [Line(line) for line in text.split("\n")]


def _docx(*paragraphs: tuple[str, int | None], table: list[list[str]] | None = None) -> bytes:
    import io

    import docx

    document = docx.Document()
    for text, level in paragraphs:
        if level is None:
            document.add_paragraph(text)
        else:
            document.add_heading(text, level=level)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(lines: list[str] | None) -> bytes:
    """A one-page PDF. ``None`` draws a line and no text, like a scanned page without OCR."""
    if lines is None:
        content = "0 0 m 200 200 l S"
    else:
        content = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(body)
    body += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1")
    body += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode("latin-1")
    body += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("latin-1")
    return body
